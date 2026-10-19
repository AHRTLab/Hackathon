#!/usr/bin/env python3
"""Builds AWARE-AI_Hackathon_Kickoff.pptx in the look of last term's BioRadio kickoff.

    pip install python-pptx
    python make_kickoff.py

It opens the BioRadio kickoff only for its theme (fonts, master), drops its slides and draws every
slide from shapes, as that deck does. Screenshots come from assets/ (taken from the live sim). Edit the
generated .pptx directly for one-off changes; edit this script to regenerate next year's deck.
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
BASE = HERE.parents[1] / "2026-bioradio/presentations/Hackathon_2026_Kickoff.pptx"
OUT = HERE / "AWARE-AI_Hackathon_Kickoff.pptx"
ASSETS = HERE / "assets"

CYAN, DCYAN, NAVY, DNAVY = "00B4C8", "008CA0", "142850", "1E3C6E"
GREY, INK, BG, LIGHT = "646E82", "1E1E32", "F0F4F8", "E6F8FA"
TRACK, STEEL, CORN, CREAM, ORANGE = "DCE1E6", "4682B4", "6495ED", "FFFAEB", "FFA500"
CONNECT, CARD_LINE, ROW = "B4DCE6", "C8D7E6", "EBF5FA"
FONT = "Calibri"


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


class Deck:
    def __init__(self) -> None:
        self.prs = Presentation(str(BASE))
        ids = self.prs.slides._sldIdLst
        for sid in list(ids):                       # keep the theme, drop last term's slides
            self.prs.part.drop_rel(sid.rId)
            ids.remove(sid)
        self.blank = next(l for l in self.prs.slide_layouts if l.name == "BLANK")
        self.slides: list = []

    # ------------------------------------------------------------------ primitives
    def new(self, bg: str):
        s = self.prs.slides.add_slide(self.blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(bg)
        self.slides.append(s)
        return s

    @staticmethod
    def rect(s, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE, line=None):
        sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
        if line:
            sh.line.color.rgb = rgb(line)
            sh.line.width = Pt(1)
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            sh.adjustments[0] = 0.08
        return sh

    @staticmethod
    def text(s, x, y, w, h, paras, size=15, color=NAVY, bold=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, spacing=0, shape=None):
        """paras: a str, or a list of paragraphs; a paragraph is a str or a list of (text, style) runs."""
        box = shape if shape is not None else s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, m, Inches(0.04))
        if isinstance(paras, str):
            paras = [paras]
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            if spacing:
                p.space_after = Pt(spacing)
            runs = [(para, {})] if isinstance(para, str) else para
            for t, st in runs:
                r = p.add_run()
                r.text = t
                f = r.font
                f.name = FONT
                f.size = Pt(st.get("size", size))
                f.bold = st.get("bold", bold)
                f.color.rgb = rgb(st.get("color", color))
                if st.get("mono"):
                    f.name = "Menlo"
        return box

    def badge(self, s, x, y, d, label, fill=CYAN, size=17, shape=MSO_SHAPE.OVAL):
        b = self.rect(s, x, y, d, d, fill, shape)
        self.text(s, 0, 0, 0, 0, label, size=size, color="FFFFFF", bold=True, align=PP_ALIGN.CENTER,
                  anchor=MSO_ANCHOR.MIDDLE, shape=b)
        return b

    # ------------------------------------------------------------------ slide frames
    def title_frame(self):
        s = self.new(NAVY)
        self.rect(s, 0, 0, 10, 0.07, CYAN)
        self.rect(s, 0, 1.5, 0.26, 2.62, CYAN)
        self.rect(s, 0, 5.55, 10, 0.07, CYAN)
        return s

    def content(self, title: str):
        s = self.new(BG)
        self.rect(s, 0, 0, 0.08, 5.625, CYAN)
        self.rect(s, 0, 0.94, 10, 0.03, CYAN)
        self.rect(s, 0, 5.59, 10, 0.04, NAVY)
        self.text(s, 0.45, 0.26, 8.63, 0.68, title, size=26, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        return s

    def numbered(self, s, items, top=1.27, gap=0.79, d=0.41, size=16, x=0.6):
        """Numbered steps joined by a thin connector. items: [(bold head, rest)]."""
        for i, (head, rest) in enumerate(items):
            y = top + i * gap
            self.badge(s, x, y + 0.04, d, str(i + 1), size=size)
            if i < len(items) - 1:
                self.rect(s, x + d / 2 - 0.015, y + 0.04 + d, 0.03, gap - d, CONNECT)
            self.text(s, x + 0.6, y, 8.3 - x, gap - 0.1,
                      [[(head, {"bold": True}), (rest, {"bold": False, "color": GREY, "size": size - 2})]],
                      size=size, anchor=MSO_ANCHOR.MIDDLE)

    def picture(self, s, path: Path, x, y, w=None, h=None, border=CARD_LINE):
        pic = s.shapes.add_picture(str(path), Inches(x), Inches(y),
                                   Inches(w) if w else None, Inches(h) if h else None)
        pic.line.color.rgb = rgb(border)
        pic.line.width = Pt(1)
        return pic

    def save(self) -> None:
        n = len(self.slides)
        for i, s in enumerate(self.slides, 1):
            dark = i in (1, n)
            self.text(s, 8.85, 5.29, 0.9, 0.26, f"{i} / {n}", size=8, color="647896" if dark else GREY,
                      align=PP_ALIGN.RIGHT)
        self.prs.save(str(OUT))
        print(f"wrote {OUT} ({n} slides)")


def build() -> None:
    d = Deck()
    T = d.text

    # 1 — title
    s = d.title_frame()
    d.rect(s, 9.38, 0.75, 0.6, 0.06, CYAN)
    d.rect(s, 8.85, 0.9, 1.13, 0.06, DCYAN)
    T(s, 0.9, 1.45, 8.25, 1.0, "AWARE-AI Hackathon 2026", size=41, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
    T(s, 0.9, 2.4, 8.25, 0.6, "Drive a robot with your muscles", size=22, color="FFFFFF", align=PP_ALIGN.CENTER)
    d.rect(s, 3.38, 3.38, 3.23, 0.03, CYAN)
    T(s, 1.5, 3.65, 6.75, 0.45, "⏱  24-Hour Build Challenge  ·  Myo armband → Clearpath Jackal",
      size=17, color=CYAN, align=PP_ALIGN.CENTER)

    # 2 — schedule (dates TBD)
    s = d.content("Schedule")
    days = [("Thursday  ·  [DATE TBD]", [("[TIME]", "Kickoff: intro, install check, first drive")]),
            ("Friday  ·  [DATE TBD]", [("8:30 am", "Breakfast, hacking starts"), ("12:00 pm", "Lunch"),
                                      ("5:00 pm", "Hacking ends · judged runs · dinner"),
                                      ("5:00–6:30 pm", "Presentations")])]
    y = 1.2
    for day, rows in days:
        T(s, 0.6, y, 8.5, 0.4, day, size=17, color=CYAN, bold=True)
        y += 0.45
        for t, what in rows:
            T(s, 0.9, y, 1.9, 0.36, t, size=15, bold=True)
            T(s, 2.8, y, 6.2, 0.36, what, size=15, color=INK)
            y += 0.4
        y += 0.2

    # 3 — what you'll build
    s = d.content("What You'll Build")
    boxes = [("Myo Armband", "(EMG + IMU)", LIGHT), ("Your Controller", "(features · model · logic)", "FFFFFF"),
             ("Command Bridge", "((v, w) at 20 Hz)", "FFFFFF"), ("Jackal Sim", "(courses · lidar · score)", "FFFFFF")]
    for i, (head, sub, fill) in enumerate(boxes):
        x = 0.53 + i * 2.44
        d.rect(s, x, 1.45, 1.8, 1.25, fill, MSO_SHAPE.ROUNDED_RECTANGLE, line=CARD_LINE)
        T(s, x + 0.07, 1.58, 1.66, 0.6, head, size=15, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        T(s, x + 0.07, 2.2, 1.66, 0.38, sub, size=11, color=GREY, align=PP_ALIGN.CENTER)
        if i < 3:
            d.rect(s, x + 1.87, 1.96, 0.45, 0.23, CYAN, MSO_SHAPE.RIGHT_ARROW)
    T(s, 0.6, 3.15, 8.63, 1.0,
      [[("Read your muscles", {"bold": True, "color": CYAN}), (", decide a forward speed ", {}),
        ("v", {"bold": True}), (" and turn rate ", {}), ("w", {"bold": True}),
        (", and drive a Clearpath Jackal through timed courses.", {})],
       [("Then keep it drivable when the link to the robot gets ", {}),
        ("slow and lossy", {"bold": True, "color": CYAN}), (".", {})]],
      size=17, color=NAVY, align=PP_ALIGN.CENTER, spacing=8)
    T(s, 0.6, 4.45, 8.63, 0.6, "(v, w) is the whole contract: any language, any model, any idea.",
      size=13, color=GREY, align=PP_ALIGN.CENTER)

    # 4 — the sim
    s = d.content("The Jackal Simulator")
    d.picture(s, ASSETS / "sim-chase.png", 0.45, 1.15, w=5.6)
    notes = [("Runs in your browser", " · MuJoCo, real Jackal dimensions"),
             ("180-beam lidar", " · your controller can read it"),
             ("Chase · Top · FPV", " cameras"),
             ("HUD", " · speed, source, collisions, deadman"),
             ("Logs every run", " · Runs CSV + 20 Hz trace")]
    for i, (h, r) in enumerate(notes):
        y = 1.2 + i * 0.62
        d.rect(s, 6.3, y + 0.12, 0.15, 0.15, CYAN, MSO_SHAPE.DIAMOND)
        T(s, 6.55, y, 3.2, 0.6, [[(h, {"bold": True}), (r, {"color": GREY, "size": 12})]], size=14)
    T(s, 0.45, 4.42, 5.6, 0.5, "ahrtlab.github.io/Hackathon  ·  or offline: jackal-bridge --serve-site",
      size=11, color=GREY, align=PP_ALIGN.CENTER)

    # 5 — the twist: telepresence
    s = d.content("The Twist: Telepresence")
    T(s, 0.6, 1.15, 8.63, 0.75,
      "Real robots are driven over real networks. Commands arrive late, and some never arrive.",
      size=17, color=INK)
    cards = [("Latency", "0 – 1000 ms", "every command is delayed"),
             ("Packet drop", "0 – 50 %", "some commands are lost"),
             ("Workload", "NASA-RTLX", "how hard was that, really?")]
    for i, (h, big, sub) in enumerate(cards):
        x = 0.6 + i * 2.95
        d.rect(s, x, 2.0, 2.65, 1.75, "FFFFFF", MSO_SHAPE.ROUNDED_RECTANGLE, line=CYAN)
        d.rect(s, x, 2.0, 2.65, 0.45, NAVY)
        T(s, x + 0.1, 2.03, 2.45, 0.4, h, size=15, color="FFFFFF", bold=True, align=PP_ALIGN.CENTER,
          anchor=MSO_ANCHOR.MIDDLE)
        T(s, x + 0.1, 2.6, 2.45, 0.55, big, size=24, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
        T(s, x + 0.1, 3.2, 2.45, 0.4, sub, size=12, color=GREY, align=PP_ALIGN.CENTER)
    d.rect(s, 0.6, 4.1, 8.63, 0.95, CREAM, MSO_SHAPE.ROUNDED_RECTANGLE)
    T(s, 0.85, 4.15, 8.2, 0.85,
      [[("Judges will also run your controller under a degraded link. ", {"bold": True, "color": NAVY}),
        ("Test with the latency / drop sliders early. A controller that only works on a perfect "
         "connection isn't finished.", {"color": INK})]], size=14, anchor=MSO_ANCHOR.MIDDLE)

    # 6 — pipeline
    s = d.content("The Hackathon Pipeline")
    d.numbered(s, [
        ("Stream  —  ", "myo_lsl_publisher.py sends Myo_EMG + Myo_IMU to LSL (no dongle)"),
        ("Look & record  —  ", "lsl-web-viewer shows the signals and records labeled training takes"),
        ("Features & model  —  ", "RMS, MAV, zero crossings… then a classifier or a proportional map"),
        ("Controller  —  ", "copy controller_template.py, fill in my_input() and step() → (v, w)"),
        ("Drive  —  ", "through jackal-bridge into the sim: gates, collisions, score"),
    ], top=1.25, gap=0.8)

    # 7 — courses
    s = d.content("The Courses")
    courses = [("L0-sandbox", "L0 Sandbox", "free driving, a few obstacles"),
               ("L1-gates", "L1 Gates", "7 gates in order · 120 s"),
               ("L2-slalom", "L2 Slalom", "8 alternating gates · 120 s"),
               ("L3-corridor-dock", "L3 Corridor + Dock", "narrow corridor, park in the dock · 150 s")]
    for i, (cid, name, sub) in enumerate(courses):
        x, y = 0.55 + (i % 2) * 4.55, 1.12 + (i // 2) * 2.1
        d.picture(s, ASSETS / f"{cid}.png", x, y, h=1.5)
        T(s, x, y + 1.52, 4.3, 0.28, [[(name, {"bold": True}), ("  " + sub, {"color": GREY, "size": 11})]], size=13)
    T(s, 3.0, 5.25, 5.7, 0.3, "Practice on all of them; judged runs are on L1–L3.", size=11, color=GREY,
      align=PP_ALIGN.RIGHT)

    # 8 — scoring
    s = d.content("How a Run Is Scored")
    d.rect(s, 0.6, 1.3, 8.63, 1.2, "FFFFFF", MSO_SHAPE.ROUNDED_RECTANGLE, line=CYAN)
    T(s, 0.7, 1.35, 8.43, 1.1,
      [[("score", {"bold": True, "color": NAVY}), ("  =  time  +  ", {}), ("5 s", {"bold": True, "color": CYAN}),
        (" × collisions  +  ", {}), ("10 s", {"bold": True, "color": CYAN}), (" × missed gates", {})]],
      size=26, color=INK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    T(s, 0.6, 2.55, 8.63, 0.35, "lower is better", size=13, color=GREY, align=PP_ALIGN.CENTER)
    for i, (h, r) in enumerate([("The clock starts", " when the robot first moves"),
                                ("Gates in order", ": the next one glows yellow; keep the left post on your left"),
                                ("Finish", ": cross the line (L3: stop inside the dock for 1 s)"),
                                ("Time limit", ": run out and it's a DNF")]):
        y = 3.05 + i * 0.52
        d.rect(s, 0.6, y + 0.12, 0.17, 0.17, CYAN, MSO_SHAPE.DIAMOND)
        T(s, 0.98, y, 8.3, 0.5, [[(h, {"bold": True}), (r, {"bold": False, "color": GREY})]], size=15)

    # 9 — rules
    s = d.content("The Rules Your Controller Lives By")
    rules = [("Deadman: 250 ms", "No command for 250 ms → the robot stops. Send at a steady 20 Hz."),
             ("Limits", "v ≤ 2.0 m/s, w ≤ 4.0 rad/s (clamped). The robot ramps; it doesn't jump."),
             ("One driver at a time", "First sender wins; another takes over after 1 s of silence."),
             ("E-stop", "\"estop\": true latches a stop; reset clears it.")]
    for i, (h, r) in enumerate(rules):
        y = 1.25 + i * 0.98
        d.rect(s, 0.6, y, 8.63, 0.8, "FFFFFF", MSO_SHAPE.ROUNDED_RECTANGLE, line=CARD_LINE)
        d.rect(s, 0.6, y, 0.06, 0.8, CYAN)
        T(s, 0.85, y + 0.05, 2.6, 0.7, h, size=16, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        T(s, 3.45, y + 0.05, 5.65, 0.7, r, size=14, color=INK, anchor=MSO_ANCHOR.MIDDLE)

    # 10 — requirements
    s = d.content("Requirements")
    reqs = ["A person drives with the Myo (EMG and/or IMU); no keyboard or gamepad in judged runs",
            "Everything runs live, in real time, through jackal-bridge",
            "Assistance is welcome (e.g. lidar keeps you off walls), but no autopilot or pre-planned route",
            "Any language, any ML framework (the bridge speaks plain JSON)"]
    for i, r in enumerate(reqs):
        y = 1.35 + i * 0.9
        d.rect(s, 0.6, y, 8.63, 0.71, "FFFFFF", MSO_SHAPE.ROUNDED_RECTANGLE, line=CARD_LINE)
        d.rect(s, 0.6, y, 0.06, 0.71, CYAN)
        d.badge(s, 0.86, y + 0.165, 0.38, "✓", size=15)
        T(s, 1.43, y + 0.05, 7.6, 0.61, r, size=15, color=INK, anchor=MSO_ANCHOR.MIDDLE)

    # 11 — deliverables
    s = d.content("Deliverables")
    for i, (h, r) in enumerate([
            ("Judged Runs", "Your controller, driven by a teammate, on L1–L3. Best score per course counts. "
                            "Plus one run under a degraded link."),
            ("Presentation (~5 min)", "Your approach, what the signals told you, results (scores, CSVs), "
                                      "and what you learned.")]):
        y = 1.35 + i * 1.8
        d.rect(s, 0.6, y, 8.63, 1.5, "FFFFFF", MSO_SHAPE.ROUNDED_RECTANGLE, line=CYAN)
        d.badge(s, 0.9, y + 0.38, 0.68, str(i + 1), fill=NAVY, size=24)
        T(s, 1.88, y + 0.22, 7.0, 0.45, h, size=20, bold=True)
        T(s, 1.88, y + 0.7, 7.0, 0.7, r, size=14, color=GREY)

    # 12 — rubric
    s = d.content("Judging Rubric (100 pts + 10 bonus)")
    for i, (h, pts, color) in enumerate([("Driving Performance", 30, CYAN), ("Creativity & Novelty", 25, STEEL),
                                         ("Control & User Experience", 25, CORN), ("Presentation Quality", 20, DNAVY)]):
        y = 1.27 + i * 0.87
        T(s, 0.6, y, 4.5, 0.38, h, size=15, bold=True)
        d.rect(s, 0.6, y + 0.38, 5.63, 0.3, TRACK)
        d.rect(s, 0.6, y + 0.38, 5.63 * pts / 30, 0.3, color)
        T(s, 6.45, y + 0.34, 1.5, 0.38, f"{pts} pts", size=17, color=color, bold=True)
    d.rect(s, 0.6, 4.85, 8.63, 0.6, CREAM, MSO_SHAPE.ROUNDED_RECTANGLE)
    T(s, 0.9, 4.88, 8.1, 0.55,
      "★  Up to +10 bonus: degraded-link robustness, cross-user calibration, multi-modal, technical depth",
      size=14, color=ORANGE, anchor=MSO_ANCHOR.MIDDLE)

    # 13 — driving performance breakdown
    s = d.content("Scoring Breakdown: Driving Performance (30 pts)")
    T(s, 0.6, 1.15, 8.6, 0.4, "Judges look for:", size=16, color=GREY)
    for i, item in enumerate(["The sim's score on L1–L3 (time + collision and gate penalties)",
                              "Completing courses at all: a slow finish beats a DNF",
                              "Smooth, deliberate control: few deadman stops, no oscillation",
                              "Consistency: repeatable runs, not one lucky one",
                              "How much the score degrades under latency / drop"]):
        y = 1.65 + i * 0.66
        d.badge(s, 0.6, y + 0.03, 0.36, "✓", size=13)
        T(s, 1.15, y, 8.0, 0.45, item, size=16, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)

    # 14 — creativity & UX
    s = d.content("Scoring: Creativity & Control / UX")
    cols = [("Creativity & Novelty (25 pts)", CYAN,
             ["Original mapping from muscles to motion", "Proportional speed, gestures, IMU steering…",
              "Smart assistance (lidar) that keeps the human in charge", "\"Wow factor\""]),
            ("Control & UX (25 pts)", STEEL,
             ["Easy to learn: a new driver is up in minutes", "Quick calibration for a new user",
              "Clear feedback: the driver knows what it will do", "Fails safe: stops when the signal is lost"])]
    for i, (h, line, items) in enumerate(cols):
        x = 0.45 + i * 4.65
        d.rect(s, x, 1.27, 4.35, 3.75, "FFFFFF", MSO_SHAPE.ROUNDED_RECTANGLE, line=line)
        d.rect(s, x, 1.27, 4.35, 0.49, NAVY)
        T(s, x + 0.23, 1.3, 3.9, 0.43, h, size=17, color="FFFFFF", bold=True, anchor=MSO_ANCHOR.MIDDLE)
        T(s, x + 0.23, 1.95, 3.9, 2.9, [f"•  {t}" for t in items], size=15, color=INK, spacing=10)

    # 15 — bonus
    s = d.content("Bonus Points (up to +10)")
    rows = [("Bonus category", "Points"), ("Robust under a degraded link (300 ms latency, 10 % drop)", "+3"),
            ("Cross-user: a new driver, calibrated in ≤ 2 min", "+3"),
            ("Multi-modal (EMG + IMU) or shared-control assistance", "+2"), ("Exceptional technical depth", "+2")]
    tbl = s.shapes.add_table(len(rows), 2, Inches(0.6), Inches(1.3), Inches(8.63), Inches(0.5 * len(rows))).table
    tbl.columns[0].width, tbl.columns[1].width = Inches(7.13), Inches(1.5)
    for r, (a, b) in enumerate(rows):
        for c, val in enumerate((a, b)):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(NAVY if r == 0 else (ROW if r % 2 else "FFFFFF"))
            T(None, 0, 0, 0, 0, val, size=15 if r else 14, bold=(r == 0 or c == 1),
              color="FFFFFF" if r == 0 else (CYAN if c == 1 else INK),
              align=PP_ALIGN.CENTER if c == 1 else PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, shape=cell)
    d.rect(s, 0.6, 4.4, 8.63, 0.6, CREAM, MSO_SHAPE.ROUNDED_RECTANGLE)
    T(s, 0.9, 4.43, 8.1, 0.55, "★  Bonus points are awarded on top of the 100-point base score by the judges.",
      size=14, color=ORANGE, anchor=MSO_ANCHOR.MIDDLE)

    # 16 — getting started
    s = d.content("Getting Started")
    mono = {"mono": True, "bold": False, "color": INK, "size": 12}
    steps = [("Install:  ", "conda env create -f environment.yml && conda activate jackal"),
             ("Sim + bridge:  ", "jackal-bridge --serve-site --open"),
             ("First drive:  ", "python …/examples/keyboard_teleop.py"),
             ("Myo:  ", "python …/examples/myo_lsl_publisher.py  then  lsl-web-viewer --open"),
             ("Build:  ", "copy controller_template.py and make it yours")]
    for i, (h, cmd) in enumerate(steps):
        y = 1.22 + i * 0.72
        d.badge(s, 0.53, y + 0.06, 0.38, str(i + 1), size=14)
        if i < len(steps) - 1:
            d.rect(s, 0.705, y + 0.44, 0.03, 0.34, CONNECT)
        T(s, 1.13, y, 8.3, 0.5, [[(h, {"bold": True}), (cmd, mono if i < 4 else {"bold": False, "color": INK})]],
          size=15, anchor=MSO_ANCHOR.MIDDLE)
    T(s, 1.13, 4.85, 8.2, 0.35,
      [[("Everything, step by step: ", {"bold": True}), ("events/2026-aware-ai-jackal/STUDENT_GUIDE.md", mono)]],
      size=13, color=CYAN)

    # 17 — tips
    s = d.content("Tips for Success")
    tips = [("Start with the fake Myo", " — fake_myo_lsl.py lets you build the pipeline without the armband"),
            ("Record training data early", " — lsl-web-viewer saves labeled takes as CSV"),
            ("Simple first", " — two reliable commands beat six flaky ones"),
            ("Always send at 20 Hz", " — or the deadman stops you; resend the last value"),
            ("Turn the latency slider on early", " — don't discover it at judging"),
            ("Test with a teammate", " — who didn't calibrate it. Then: myo_check.py --off")]
    for i, (h, r) in enumerate(tips):
        y = 1.2 + i * 0.68
        d.rect(s, 0.6, y + 0.13, 0.17, 0.17, CYAN, MSO_SHAPE.DIAMOND)
        T(s, 0.98, y, 8.4, 0.53, [[(h, {"bold": True}), (r, {"bold": False, "color": GREY})]], size=15)

    # 18 — close
    s = d.title_frame()
    T(s, 0.75, 1.5, 8.63, 0.9, "Build Something Amazing", size=36, color="FFFFFF", bold=True, align=PP_ALIGN.CENTER)
    d.rect(s, 3.38, 2.55, 3.23, 0.03, CYAN)
    T(s, 1.13, 2.85, 7.88, 0.38, "Good luck!", size=18, color="00D2E6", align=PP_ALIGN.CENTER)
    T(s, 1.5, 3.6, 6.75, 1.0, ["Repo: github.com/AHRTLab/Hackathon", "Sim: ahrtlab.github.io/Hackathon",
                               "Contact: jrheee@rit.edu"],
      size=14, color="C8D7E6", align=PP_ALIGN.CENTER, spacing=4)

    d.save()


if __name__ == "__main__":
    build()
