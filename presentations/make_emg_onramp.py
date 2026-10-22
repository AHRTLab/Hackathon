#!/usr/bin/env python3
"""Builds EMG_On-Ramp.pptx: a short, no-background introduction to EMG and the Myo, in the kickoff's look.

    pip install python-pptx matplotlib scipy
    python emg_figures.py            # (re)draw the simulated-signal figures in assets/emg/
    python make_emg_onramp.py        # refuses to overwrite the (hand-edited) .pptx; --force to do so

The full EMG lecture from last term (electromyography.pptx) is on the `2026-bioradio` branch.
"""
from __future__ import annotations

import sys
from pathlib import Path

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches

from make_kickoff import (CARD_LINE, CREAM, CYAN, DCYAN, GREY, INK, LIGHT, NAVY, ORANGE, ROW, STEEL,
                          Deck)

HERE = Path(__file__).resolve().parent
OUT = HERE / "EMG_On-Ramp.pptx"
FIG = HERE / "assets" / "emg"


def diamonds(d: Deck, s, items, x=0.6, top=1.2, gap=0.62, w=8.6, size=15):
    for i, (h, r) in enumerate(items):
        y = top + i * gap
        d.rect(s, x, y + 0.12, 0.16, 0.16, CYAN, MSO_SHAPE.DIAMOND)
        d.text(s, x + 0.36, y, w - 0.36, gap, [[(h, {"bold": True}), (r, {"bold": False, "color": GREY})]], size=size)


def build(force: bool = False) -> None:
    d = Deck()
    T = d.text

    # 1 — title
    s = d.title_frame()
    d.rect(s, 9.38, 0.75, 0.6, 0.06, CYAN)
    d.rect(s, 8.85, 0.9, 1.13, 0.06, DCYAN)
    T(s, 0.9, 1.45, 8.25, 1.0, "EMG On-Ramp", size=41, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
    T(s, 0.9, 2.4, 8.25, 0.6, "What your Myo is actually measuring", size=22, color="FFFFFF", align=PP_ALIGN.CENTER)
    d.rect(s, 3.38, 3.38, 3.23, 0.03, CYAN)
    T(s, 1.5, 3.65, 6.75, 0.45, "⏱  15 minutes  ·  no background needed", size=17, color=CYAN, align=PP_ALIGN.CENTER)

    # 2 — why
    s = d.content("Why Understand the Signal?")
    cards = [("Your code only sees numbers", "8 streams of numbers, 200 times a second. Knowing what they mean is "
                                             "how you turn them into a good command."),
             ("Intent before motion", "Muscle activity starts ~50–100 ms before the hand actually moves. "
                                      "That is the head start EMG gives a controller."),
             ("Fewer surprises", "Noise, fatigue and a shifted band all look like “my classifier broke.” "
                                 "Recognise them and you can fix them.")]
    for i, (h, body) in enumerate(cards):
        x = 0.6 + i * 2.95
        d.rect(s, x, 1.45, 2.65, 2.75, "FFFFFF", MSO_SHAPE.ROUNDED_RECTANGLE, line=CYAN)
        d.rect(s, x, 1.45, 2.65, 0.6, NAVY)
        T(s, x + 0.1, 1.48, 2.45, 0.55, h, size=14, color="FFFFFF", bold=True, align=PP_ALIGN.CENTER,
          anchor=MSO_ANCHOR.MIDDLE)
        T(s, x + 0.2, 2.25, 2.25, 1.9, body, size=15, color=INK)

    # 3 — where it comes from
    s = d.content("Where the Signal Comes From")
    chain = [("Brain", "decides to move"), ("Motor neurons", "fire spikes"), ("Muscle fibres", "contract"),
             ("Electric field", "reaches the skin"), ("Myo", "8 electrodes")]
    for i, (h, sub) in enumerate(chain):
        x = 0.45 + i * 1.88
        d.rect(s, x, 1.3, 1.5, 1.05, LIGHT if i in (0, 4) else "FFFFFF", MSO_SHAPE.ROUNDED_RECTANGLE, line=CARD_LINE)
        T(s, x + 0.05, 1.38, 1.4, 0.45, h, size=14, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        T(s, x + 0.05, 1.83, 1.4, 0.4, sub, size=11, color=GREY, align=PP_ALIGN.CENTER)
        if i < 4:
            d.rect(s, x + 1.55, 1.72, 0.28, 0.2, CYAN, MSO_SHAPE.RIGHT_ARROW)
    diamonds(d, s, [("Motor unit", " = one motor neuron + all the muscle fibres it controls. It fires all-or-none."),
                    ("More force", " = more motor units switched on, each firing faster."),
                    ("Surface EMG", " = the sum of all those tiny electrical spikes, seen through skin and fat."),
                    ("So", ": it is a crowd's noise, not a single voice. Its loudness is the information.")],
             top=2.75, gap=0.6, size=14)

    # 4 — effort
    s = d.content("More Effort, Bigger Signal")
    d.picture(s, FIG / "effort.png", 0.6, 1.15, w=8.6, border=None)
    T(s, 0.6, 4.25, 8.6, 0.9,
      [[("Amplitude tracks effort. ", {"bold": True}),
        ("A light squeeze, a medium one and a hard one: the signal gets larger and denser as more motor units join in.",
         {"color": INK})],
       [("Simulated signal at the Myo's 200 Hz, for illustration.", {"size": 10, "color": GREY})]],
      size=15, spacing=4)

    # 5 — what raw EMG looks like
    s = d.content("Reading Raw EMG")
    diamonds(d, s, [("It looks like noise", " — and statistically it is: random-looking, centred on zero."),
                    ("The sign means nothing", " — +80 and −80 are the same amount of activity."),
                    ("One sample tells you nothing", " — always look at a window of samples (e.g. the last 200 ms)."),
                    ("Loudness is the information", " — how big the wiggles are over a window = how hard the muscle works."),
                    ("Where it is loud is the information too", " — which of the 8 channels light up = which muscles = which gesture.")],
             top=1.25, gap=0.72, size=15)

    # 6 — the Myo
    s = d.content("Meet the Myo Armband")
    specs = [("8 EMG channels", " · electrode pods around the forearm"),
             ("200 Hz", " · 200 samples per second per channel"),
             ("8-bit values", " · −128 … 127, unitless (not µV)"),
             ("Dry electrodes", " · no gel: needs a snug fit and ~2 min to settle"),
             ("Plus an IMU", " · orientation, acceleration, rotation at 50 Hz"),
             ("On LSL", " · Myo_EMG, Myo_IMU, Myo_Events (arm sync)")]
    diamonds(d, s, specs, top=1.2, gap=0.6, w=5.2, size=14)
    import math
    cx, cy, r = 7.55, 2.95, 1.25
    d.rect(s, cx - r - 0.32, cy - r - 0.32, 2 * r + 0.64, 2 * r + 0.64, "FFFFFF", MSO_SHAPE.OVAL, line=CARD_LINE)
    d.rect(s, cx - 0.55, cy - 0.4, 1.1, 0.8, LIGHT, MSO_SHAPE.OVAL)
    T(s, cx - 0.6, cy - 0.25, 1.2, 0.5, "forearm\n(cross-section)", size=9, color=GREY, align=PP_ALIGN.CENTER,
      anchor=MSO_ANCHOR.MIDDLE)
    for k in range(8):
        a = math.pi / 2 - k * 2 * math.pi / 8
        d.badge(s, cx + r * math.cos(a) - 0.21, cy - r * math.sin(a) - 0.21, 0.42, str(k + 1),
                fill=CYAN if k in (0, 1, 2, 7) else STEEL, size=13, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    T(s, 5.9, 4.55, 3.4, 0.3, [[("■ ", {"color": CYAN}), ("palm side (flexors)   ", {}),
                               ("■ ", {"color": STEEL}), ("back (extensors)", {})]], size=10, color=GREY,
      align=PP_ALIGN.CENTER)
    T(s, 5.9, 4.82, 3.4, 0.3, "which pod sits where depends on how the band is worn", size=10, color=GREY,
      align=PP_ALIGN.CENTER)

    # 7 — gestures
    s = d.content("Different Gestures, Different Channels")
    d.picture(s, FIG / "gestures.png", 0.5, 1.15, w=8.9, border=None)
    diamonds(d, s, [("A gesture is a pattern", " across the 8 channels, not a single number."),
                    ("Rotate the band and the pattern rotates too", ": calibrate for each person and each fitting.")],
             top=3.65, gap=0.55, size=14)
    T(s, 0.6, 4.85, 8.6, 0.3, "Illustrative patterns, not measured data. See your own in lsl-web-viewer.",
      size=10, color=GREY)

    # 8 — processing
    s = d.content("From Raw Signal to Envelope")
    d.picture(s, FIG / "pipeline.png", 0.45, 1.08, h=4.1, border=None)
    steps = [("Raw", "what the Myo sends"), ("Remove offset", "subtract the mean, so rest sits at 0"),
             ("Rectify", "take |x|: negative wiggles count as activity too"),
             ("Smooth", "low-pass or a moving average: a slow “how hard” curve")]
    for i, (h, r_) in enumerate(steps):
        y = 1.2 + i * 0.95
        d.badge(s, 7.35, y, 0.36, str(i + 1), size=13)
        T(s, 7.8, y - 0.05, 2.0, 0.35, h, size=14, bold=True)
        T(s, 7.8, y + 0.28, 2.0, 0.6, r_, size=11, color=GREY)
    T(s, 7.35, 5.0, 2.5, 0.3, "That curve can drive speed directly.", size=11, color=CYAN, bold=True)

    # 9 — windows & features
    s = d.content("Windows and Features")
    d.picture(s, FIG / "windows.png", 0.6, 1.08, w=5.2, border=None)
    rows = [("Feature", "Tells you"), ("RMS / MAV", "how hard (amplitude)"), ("WL  waveform length", "how busy the signal is"),
            ("ZC  zero crossings", "a rough frequency"), ("SSC  slope sign changes", "high-frequency content")]
    tbl = s.shapes.add_table(len(rows), 2, Inches(6.0), Inches(1.15), Inches(3.6), Inches(0.42 * len(rows))).table
    tbl.columns[0].width, tbl.columns[1].width = Inches(1.85), Inches(1.75)
    for ri, (a, b) in enumerate(rows):
        for ci, val in enumerate((a, b)):
            cell = tbl.cell(ri, ci)
            cell.fill.solid()
            from pptx.dml.color import RGBColor
            cell.fill.fore_color.rgb = RGBColor.from_string(NAVY if ri == 0 else (ROW if ri % 2 else "FFFFFF"))
            T(None, 0, 0, 0, 0, val, size=11, bold=(ri == 0 or ci == 0), color="FFFFFF" if ri == 0 else INK,
              anchor=MSO_ANCHOR.MIDDLE, shape=cell)
    d.rect(s, 0.6, 3.55, 9.0, 1.45, "FFFFFF", MSO_SHAPE.ROUNDED_RECTANGLE, line=CYAN)
    T(s, 0.85, 3.62, 8.55, 1.35,
      [[("Every 50 ms: ", {"bold": True}), ("take the last 200 ms (40 samples) of each channel, compute a few features.", {})],
       [("8 channels × 4 features = ", {"bold": True}), ("32 numbers", {"bold": True, "color": CYAN}),
        (" describing “what the arm is doing right now.” That vector is what a classifier sees.", {})],
       [("At 200 Hz the Myo only captures up to 100 Hz, so amplitude features (RMS, MAV, WL) work best.",
         {"size": 12, "color": GREY})]],
      size=14, color=INK, spacing=5, anchor=MSO_ANCHOR.MIDDLE)

    # 10 — two ways to control
    s = d.content("Two Ways to Turn EMG Into (v, w)")
    cols = [("Proportional", CYAN, ["Envelope → speed directly", "Smooth and fast to build",
                                    "Squeeze harder = go faster", "Needs per-person calibration",
                                    "e.g. myo_emg_example.py"]),
            ("Gesture classification", STEEL, ["Features → classifier → a class", "fist / wave in / wave out / rest",
                                              "Each class = a command", "Needs labeled training data",
                                              "LDA or SVM is a great start"])]
    for i, (h, line, items) in enumerate(cols):
        x = 0.45 + i * 4.65
        d.rect(s, x, 1.2, 4.35, 2.95, "FFFFFF", MSO_SHAPE.ROUNDED_RECTANGLE, line=line)
        d.rect(s, x, 1.2, 4.35, 0.49, NAVY)
        T(s, x + 0.23, 1.23, 3.9, 0.43, h, size=17, color="FFFFFF", bold=True, anchor=MSO_ANCHOR.MIDDLE)
        T(s, x + 0.23, 1.85, 3.9, 2.2, [f"•  {t}" for t in items], size=14, color=INK, spacing=6)
    d.rect(s, 0.45, 4.35, 9.0, 0.7, CREAM, MSO_SHAPE.ROUNDED_RECTANGLE)
    T(s, 0.7, 4.38, 8.5, 0.64,
      [[("★  Mix them: ", {"bold": True}), ("the class picks the direction, the envelope picks the speed.", {})]],
      size=15, color=ORANGE, anchor=MSO_ANCHOR.MIDDLE)

    # 11 — calibration
    s = d.content("Calibration: Every Arm Is Different")
    d.numbered(s, [("Rest  —  ", "relax for a few seconds: each channel's baseline"),
                   ("Max  —  ", "a firm (not painful) contraction for each gesture you use"),
                   ("Normalise  —  ", "(x − rest) / (max − rest) → 0 … 1 per channel"),
                   ("Deadband  —  ", "treat < ~15 % as zero, so resting noise doesn't drive the robot"),
                   ("Redo it  —  ", "new person, band moved, or after a break")], top=1.2, gap=0.79, size=15)

    # 12 — what goes wrong
    s = d.content("What Goes Wrong (and the Fix)")
    rows = [("You see", "Likely cause", "Fix"),
            ("Big noise on every channel", "loose band, dry skin", "tighten (sizing clips); wait 2 min"),
            ("Spikes when the arm moves", "motion artefacts", "steady arm; filter; ignore short spikes"),
            ("Wrong gesture detected", "band rotated since calibration", "re-sync and recalibrate"),
            ("Signal creeps up over time", "fatigue", "rest breaks; recalibrate"),
            ("Two gestures look the same", "crosstalk between muscles", "more distinct gestures; better features"),
            ("Robot creeps at rest", "rest isn't zero", "deadband after calibration")]
    tbl = s.shapes.add_table(len(rows), 3, Inches(0.5), Inches(1.15), Inches(9.0), Inches(0.5 * len(rows))).table
    for ci, w in enumerate((3.0, 2.8, 3.2)):
        tbl.columns[ci].width = Inches(w)
    from pptx.dml.color import RGBColor
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(NAVY if ri == 0 else (ROW if ri % 2 else "FFFFFF"))
            T(None, 0, 0, 0, 0, val, size=13 if ri else 13, bold=(ri == 0 or ci == 0),
              color="FFFFFF" if ri == 0 else (NAVY if ci == 0 else INK), anchor=MSO_ANCHOR.MIDDLE, shape=cell)

    # 13 — fatigue
    s = d.content("Fatigue Changes the Signal")
    d.picture(s, FIG / "fatigue.png", 0.5, 1.2, w=4.6, border=None)
    diamonds(d, s, [("Same force, bigger signal", ": tired muscles recruit more units"),
                    ("Lower frequencies", ": the spectrum shifts down (median frequency drops)"),
                    ("Your thresholds drift", ": what was “medium” at 9 am is “hard” at 3 pm"),
                    ("Plan for it", ": short breaks, quick recalibration, relative (normalised) features")],
             x=5.3, top=1.3, gap=0.85, w=4.4, size=14)

    # 14 — try it
    s = d.content("Try It Now (10 minutes)")
    d.picture(s, FIG / "viewer-emg.png", 5.05, 1.15, w=4.6)
    hands = [("Stream: ", "python examples/myo_lsl_publisher.py (or fake_myo_lsl.py)"),
             ("Look: ", "lsl-web-viewer --open, tick Myo_EMG"),
             ("Play: ", "relax, fist, wave in, wave out. Which bars move?"),
             ("Record: ", "label a few 5 s takes of each gesture"),
             ("Drive: ", "python examples/myo_emg_example.py")]
    for i, (h, r_) in enumerate(hands):
        y = 1.2 + i * 0.75
        d.badge(s, 0.53, y + 0.05, 0.36, str(i + 1), size=13)
        if i < len(hands) - 1:
            d.rect(s, 0.695, y + 0.41, 0.03, 0.39, "B4DCE6")
        T(s, 1.05, y, 3.9, 0.7, [[(h, {"bold": True}), (r_, {"bold": False, "color": INK, "size": 12})]], size=14,
          anchor=MSO_ANCHOR.MIDDLE)
    T(s, 5.05, 3.5, 4.6, 0.3, "lsl-web-viewer: live traces + RMS bars (fake Myo shown)", size=10, color=GREY,
      align=PP_ALIGN.CENTER)

    # 15 — cheat sheet
    s = d.content("Numbers to Remember")
    nums = [("8", "channels"), ("200 Hz", "sample rate"), ("±128", "value range"),
            ("200 ms", "feature window"), ("20 Hz", "send commands"), ("250 ms", "deadman stop")]
    for i, (big, lab) in enumerate(nums):
        x, y = 0.6 + (i % 3) * 3.0, 1.3 + (i // 3) * 1.85
        d.rect(s, x, y, 2.7, 1.55, "FFFFFF", MSO_SHAPE.ROUNDED_RECTANGLE, line=CARD_LINE)
        T(s, x, y + 0.15, 2.7, 0.8, big, size=34, color=CYAN, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        T(s, x, y + 0.95, 2.7, 0.45, lab, size=14, color=GREY, align=PP_ALIGN.CENTER)

    # 16 — close
    s = d.title_frame()
    T(s, 0.75, 1.5, 8.63, 0.9, "Questions?", size=36, color="FFFFFF", bold=True, align=PP_ALIGN.CENTER)
    d.rect(s, 3.38, 2.55, 3.23, 0.03, CYAN)
    T(s, 1.13, 2.85, 7.88, 0.38, "Go look at your own muscles.", size=18, color="00D2E6", align=PP_ALIGN.CENTER)
    T(s, 1.0, 3.55, 8.0, 1.3, ["Setup: STUDENT_GUIDE.md  ·  Smallest controller: examples/myo_emg_example.py",
                               "Deeper dive: electromyography.pptx on the 2026-bioradio branch",
                               "github.com/AHRTLab/Hackathon"],
      size=13, color="C8D7E6", align=PP_ALIGN.CENTER, spacing=4)

    d.save(OUT, force=force)


if __name__ == "__main__":
    build(force="--force" in sys.argv)
